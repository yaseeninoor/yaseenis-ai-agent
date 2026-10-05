import os, hashlib, pathlib, requests, csv
from openpyxl import load_workbook
from pypdf import PdfReader
from docx import Document

ROOT=pathlib.Path("FM-STORE-SETUP/source-files")
URL=os.environ["INGEST_URL"]; TOKEN=os.environ["INGEST_TOKEN"]
REPO=os.environ["REPO"]; REF=os.environ["REF"]
EXT={".xlsx",".xlsm",".pdf",".docx",".csv",".txt",".md"}

def chunks(text,n=3500):
    text=" ".join(str(text or "").split())
    return [text[i:i+n] for i in range(0,len(text),n)] or []

def emit(path,text,sheet=None,page=None):
    for i,c in enumerate(chunks(text)):
        h=hashlib.sha256(f"{path}|{sheet}|{page}|{i}|{c}".encode()).hexdigest()
        yield {"file_path":str(path),"sheet_name":sheet,"page_no":page,"chunk_no":i,"content":c,"content_hash":h}

out=[]
for p in ROOT.rglob("*"):
    if not p.is_file() or p.suffix.lower() not in EXT: continue
    rel=p.relative_to(pathlib.Path("."))
    ext=p.suffix.lower()
    try:
        if ext in {".xlsx",".xlsm"}:
            wb=load_workbook(p,read_only=True,data_only=True)
            for s in wb.sheetnames:
                ws=wb[s]; lines=[]
                for row in ws.iter_rows(values_only=True):
                    vals=[str(v) if v is not None else "" for v in row]
                    if any(vals): lines.append(" | ".join(vals))
                out += list(emit(rel,"\n".join(lines),s))
        elif ext==".pdf":
            r=PdfReader(str(p))
            for i,page in enumerate(r.pages,1):
                out += list(emit(rel,page.extract_text() or "",page=i))
        elif ext==".docx":
            d=Document(str(p)); text="\n".join(x.text for x in d.paragraphs)
            for t in d.tables:
                for row in t.rows: text += "\n"+" | ".join(c.text for c in row.cells)
            out += list(emit(rel,text))
        elif ext==".csv":
            with open(p,encoding="utf-8-sig",newline="") as f: out += list(emit(rel,"\n".join(" | ".join(r) for r in csv.reader(f))))
        else:
            out += list(emit(rel,p.read_text(encoding="utf-8",errors="ignore")))
    except Exception as e:
        print("SKIP",rel,e)

for i in range(0,len(out),50):
    r=requests.post(URL,headers={"Authorization":f"Bearer {TOKEN}","Content-Type":"application/json"},
                    json={"repo":REPO,"ref":REF,"chunks":out[i:i+50]},timeout=60)
    r.raise_for_status()
print("Indexed",len(out),"chunks")
